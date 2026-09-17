"""The Excel export must carry the pipeline's numbers and stay a valid, live workbook.

These tests open the produced file with openpyxl, so they check structure and formulas, not what
Excel renders. The formulas themselves were verified once against a real Excel recalculation.
"""
import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from core import export_excel as XL

pytestmark = pytest.mark.filterwarnings("ignore::UserWarning")


def _opt(res, **kw):
    return XL.ExportOptions(wells=tuple(res.wells_all),
                            start=str(res.meta["rt_start"].date()),
                            end=str(res.meta["rt_end"].date()), **kw)


@pytest.fixture(scope="module")
def book(res):
    """The default export of GC31, loaded back with its formulas intact."""
    b = XL.workbook_bytes(res, _opt(res))
    return load_workbook(io.BytesIO(b))


def test_one_tab_per_analysed_well(book, res):
    for w in res.wells_analysed:
        assert w in book.sheetnames
    # an excluded well is reported in the tables but has no tab: it carries no K to chart
    for w in res.excluded["WELL_NAME"]:
        assert w not in book.sheetnames


def test_every_table_sheet_is_present(book):
    for name in ["Read me", "Wells overview", "Matched tests", "Calibration K", "All well tests",
                 "Validation", "Error by method", "Data quality", "Events", "Wells and pumps"]:
        assert name in book.sheetnames


def test_tables_and_charts_exist(book, res):
    for w in res.wells_analysed:
        ws = book[w]
        assert len(ws.tables) >= 3, f"{w} lost its data tables"
        assert len(ws._charts) >= 5, f"{w} lost its charts"
    assert len(book["Wells overview"]._charts) == 2


def test_table_headers_are_unique_per_table(book):
    """Excel refuses to open a workbook whose Table repeats a header name."""
    for ws in book.worksheets:
        for t in ws.tables.values():
            names = [c.name for c in t.tableColumns]
            assert len(names) == len(set(names)), f"{ws.title}/{t.displayName}: {names}"


def test_k_is_live_from_the_matched_tests(book, res):
    """K single is a MEDIAN over the well's block of non-suspect matched tests."""
    cal = book["Calibration K"]
    row = next(r for r in cal.iter_rows(min_row=1, max_row=cal.max_row)
               if r[0].value == res.wells_analysed[0])
    k_cell = next(c for c in row if isinstance(c.value, str) and c.value.startswith("=IFERROR(MEDIAN("))
    assert "Matched tests" in k_cell.value
    # ... and the well tab points at that cell rather than hard-coding the number
    ws = book[res.wells_analysed[0]]
    assert ws["B7"].value.startswith("='Calibration K'!")


def test_well_tab_rate_follows_the_k_cell(book, res):
    w = res.wells_analysed[0]
    ws = book[w]
    name = XL._table_name("K", w)
    assert name in book.defined_names
    first = ws.cell(row=XL.DATA_ROW + 1, column=3)          # the 'Q at K cell' column
    assert name in str(first.value) and "X" not in str(first.value).split("*")[0]


def test_values_match_the_pipeline(book, res):
    """The exported series is the app's own numbers, not a re-derivation."""
    w = res.wells_analysed[0]
    opt = _opt(res)
    s = XL.well_series(res, w, opt)
    ws = book[w]
    assert ws.cell(row=XL.DATA_ROW, column=1).value == "Time"
    got = ws.cell(row=XL.DATA_ROW + 1, column=2).value      # Q virtual
    assert got == pytest.approx(float(s["Q_app"].iloc[0]), rel=1e-9)


def test_values_only_export_has_no_formulas(res):
    b = XL.workbook_bytes(res, _opt(res, live_formulas=False, per_well_tabs=True))
    wb = load_workbook(io.BytesIO(b))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                assert not (isinstance(c.value, str) and c.value.startswith("=")), \
                    f"{ws.title}!{c.coordinate} is still a formula"


def test_water_cut_export_uses_the_downhole_k(res):
    b = XL.workbook_bytes(res, _opt(res, wc_correction=True, k_mode="single", freq="h"))
    wb = load_workbook(io.BytesIO(b))
    ws = wb[res.wells_analysed[0]]
    assert ws.cell(row=7, column=1).value == "K_dh for this well"
    q = str(ws.cell(row=XL.DATA_ROW + 1, column=3).value)
    assert "/" in q, "the corrected rate must divide by B_liq"


def test_meleiha_exports_without_a_bubble_point(mel):
    """The second field has several electrical regimes and no laboratory PVT, so the free-gas
    tables are empty. The workbook must still build, with a charted tab per analysed well."""
    assert not len(mel.gas) and not len(mel.lab_pvt)
    b = XL.workbook_bytes(mel, _opt(mel))
    wb = load_workbook(io.BytesIO(b))
    assert set(mel.wells_analysed) <= set(wb.sheetnames)
    for w in mel.wells_analysed:
        assert len(wb[w]._charts) >= 5
    # a well split into regimes gets one calibration row per regime, all of them live
    cal = wb["Calibration K"]
    wells = [cal.cell(row=r, column=1).value for r in range(1, cal.max_row + 1)]
    assert len(mel.cal) == sum(1 for w in wells if w in set(mel.cal["WELL_NAME"]))


def test_period_is_respected(res):
    start, end = "2025-01-01", "2025-03-31"
    b = XL.workbook_bytes(res, XL.ExportOptions(wells=tuple(res.wells_analysed),
                                                start=start, end=end))
    wb = load_workbook(io.BytesIO(b))
    ws = wb[res.wells_analysed[0]]
    times = [ws.cell(row=r, column=1).value
             for r in range(XL.DATA_ROW + 1, min(ws.max_row, XL.DATA_ROW + 200) + 1)]
    times = [t for t in times if isinstance(t, (pd.Timestamp, type(pd.Timestamp.now().to_pydatetime())))]
    assert times, "no rows exported for the period"
    assert min(times) >= pd.Timestamp(start)
    assert max(times) < pd.Timestamp(end) + pd.Timedelta(days=1)
